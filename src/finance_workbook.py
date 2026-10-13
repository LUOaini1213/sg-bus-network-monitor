"""Write outputs/bus_finance_model.xlsx: inputs in blue, every derived cell a live Excel formula.

Sheets: README, Inputs, Fares, Rides, Revenue, Transfers, Reconciliation, Fund, Contracts, Projection, Scenarios,
Sources. Run after finance.py (python src/finance_workbook.py); finance_model.py holds the same model in Python.
"""
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

import finance_model as fm

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
XLSX = OUT / "bus_finance_model.xlsx"
BLUE = Font(color="1F4E9E")
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="E8EDF4")
N = fm.N_BANDS
GROUP_LABEL = {"adult": "Adult", "senior": "Senior / PwD", "student": "Student", "workfare": "Workfare (WTCS)"}


def name(wb, label, ref):
    wb.defined_names[label] = DefinedName(label, attr_text=ref)


def header(ws, row, labels, widths=None):
    for j, t in enumerate(labels, 1):
        c = ws.cell(row, j, t)
        c.font = BOLD
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[L(j)].width = w


def put_input(ws, row, label, value, unit, source, wb, nm, fmt=None):
    ws.cell(row, 1, label)
    c = ws.cell(row, 2, value)
    c.font = BLUE
    if fmt:
        c.number_format = fmt
    ws.cell(row, 3, unit)
    ws.cell(row, 4, source)
    name(wb, nm, f"Inputs!$B${row}")


def build(rides_csv=None, assumptions=None, sources=None, stats=None, path=XLSX):
    sources = sources or fm.load_sources()
    a = assumptions or json.loads(fm.PARAMS.read_text(encoding="utf-8"))
    t = pd.read_csv(rides_csv or OUT / "finance_rides_by_band.csv")
    stats = stats or json.loads((OUT / "finance_inputs.json").read_text(encoding="utf-8"))
    cur, new = fm.load_fares(fm.CURRENT, sources), fm.load_fares(fm.APPROVED, sources)
    fund = [r for r in sources["revenue_cost"] if str(r.get("entity", "")).startswith("LTA Bus Contracting")]
    fund = sorted(fund, key=lambda r: r["period"])
    wb = Workbook()

    # ---------------------------------------------------------------- README
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 120
    lines = [
        "Singapore bus finance and fare model (public data only)",
        "",
        "Blue cells are inputs; every other number is a live formula. Change an input and the model recalculates.",
        "Rides: DataMall origin-destination volumes, July and August 2026, scaled to 2026 day counts (Rides sheet).",
        "Fares: PTC card fare tables in force from 27 Dec 2025 and approved for 26 Dec 2026 (Fares sheet).",
        "Published figures: LTA audited Bus Contracting accounts, LTA ridership statistics, PTC fare reviews, LTA "
        "contract award releases (Sources sheet, with the quoted text).",
        "Fare-group shares and the transfer share are derived from published counts under stated assumptions "
        "(Inputs sheet).",
        "Monthly passes, cash fares (<1% of journeys), express surcharges on mixed pairs and fare evasion are not "
        "modelled separately.",
        "Built by src/finance_workbook.py; src/finance_model.py computes the same headline cells in Python and the "
        "test suite checks that the two agree.",
    ]
    for i, s in enumerate(lines, 1):
        ws.cell(i, 1, s).font = BOLD if i == 1 else Font()

    # ---------------------------------------------------------------- Inputs
    ws = wb.create_sheet("Inputs")
    header(ws, 1, ["Input", "Value", "Unit", "Source / note"], [48, 16, 16, 110])
    r = 2
    d = a["derivation"]
    raw = [("Bus rides a day on senior concession cards, 2019", d["senior_bus_trips_2019"], "rides/day",
            d["senior_note"], "senior_trips_2019", "#,##0"),
           ("LTA average daily public bus ridership, 2019", d["bus_rides_2019"], "rides/day", d["ridership_note"],
            "bus_rides_2019", "#,##0")]
    for g in ["senior", "pwd", "student", "workfare", "child"]:
        raw.append((f"Share of concession cards: {g}", d["card_shares"][g], "share", d["cards_note"], f"card_{g}",
                    "0%"))
    raw += [("Smart-card trips on a weekday (bus + rail, Feb 2018)", d["trips_2018"], "trips/day",
             d["journeys_note"], "trips_2018", "#,##0"),
            ("Linked journeys on that weekday", d["journeys_2018"], "journeys/day", d["journeys_note"],
             "journeys_2018", "#,##0")]
    for lab, val, unit, note, nm, fmt in raw:
        put_input(ws, r, lab, val, unit, note, wb, nm, fmt)
        r += 1
    ws.cell(r, 1, "Derived shares of bus rides (formulas). Assumes rides per card are the same in every concession "
                  "group, calibrated on seniors; children under 7 ride free.").font = BOLD
    r += 1
    derived = {
        "senior": "=senior_trips_2019/bus_rides_2019*(card_senior+card_pwd)/card_senior",
        "student": "=senior_trips_2019/bus_rides_2019*card_student/card_senior",
        "workfare": "=senior_trips_2019/bus_rides_2019*card_workfare/card_senior",
        "free": "=senior_trips_2019/bus_rides_2019*card_child/card_senior",
    }
    for g in ["senior", "student", "workfare", "free"]:
        label = (f"Share of rides paying {GROUP_LABEL[g]} fares" if g != "free"
                 else "Share of rides that are free (children under 7)")
        ws.cell(r, 1, label)
        ws.cell(r, 2, derived[g]).number_format = "0.0%"
        name(wb, f"share_{g}", f"Inputs!$B${r}")
        r += 1
    ws.cell(r, 1, "Share of rides paying Adult fares (the rest)")
    ws.cell(r, 2, "=1-share_senior-share_student-share_workfare-share_free").number_format = "0.0%"
    name(wb, "share_adult", f"Inputs!$B${r}")
    r += 1
    ws.cell(r, 1, "Share of rides that are a transfer leg (1 - journeys / trips)")
    ws.cell(r, 2, "=1-journeys_2018/trips_2018").number_format = "0.0%"
    ws.cell(r, 4, "Bus and rail combined; used for bus rides as an approximation")
    name(wb, "transfer_share", f"Inputs!$B${r}")
    r += 2
    rows = [
        ("20th percentile monthly household market income, 2025", a["income_p20"], "S$/month", a["income_note"],
         "income_p20", "#,##0"),
        ("LTA average daily public bus ridership, 2025", a["lta_avg_daily_bus_ridership_2025"], "rides/day",
         a["lta_ridership_note"], "lta_daily", "#,##0"),
        ("FRE 2025 extra fare revenue, bus + TEL, per year", a["fre2025_bus_tel_increase"], "S$/yr",
         a["fre2025_note"], "fre2025_inc", "#,##0"),
        ("FRE 2026 extra fare revenue, bus + TEL, per year", a["fre2026_bus_tel_increase"], "S$/yr",
         a["fre2026_note"], "fre2026_inc", "#,##0"),
        ("Days of FY2025/26 at the old (Dec 2024) fares", a["fy25_days_at_old_fares"], "days",
         "1 Apr 2025 to 26 Dec 2025; the 2025 fare review took effect on 27 Dec 2025", "old_days", "0"),
        ("Rail contracting (TEL/JRL) fare revenue, FY2025/26", a["tel_fare_revenue_fy25"], "S$",
         a["tel_note"], "tel_rev", "#,##0"),
        ("Scheduled weekday bus-km per day (DataMall timetable bands)", stats["weekday_bus_km"], "bus-km/day",
         "finance.py: route length x scheduled trips from headway bands and first/last bus; weekday timetable",
         "wd_bus_km", "#,##0"),
        ("Rides a month for a regular commuter", a["commuter_rides_per_month"], "rides",
         "2 rides a working day x 22 days (illustrative)", "commuter_rides", "0"),
        ("Adult Monthly Travel Pass (bus + train), from 27 Dec 2025", a["adult_pass"], "S$/month",
         "PTC FRE 2025 Annex B", "adult_pass", "0.00"),
        ("Ridership growth per year (projection)", a["ridership_growth"], "share", a["ridership_growth_note"],
         "g_rides", "0.00%"),
        ("Operating cost growth per year (projection)", a["cost_growth"], "share", a["cost_growth_note"],
         "g_cost", "0.00%"),
        ("Fare change at each later fare review (projection, FY2027 on)", a["later_fre"], "share",
         "Not decided; 0% by default. Fare reviews take effect in late December.", "later_fre", "0.0%"),
        ("Concession scenario: change in senior / PwD fares", a["concession_change"], "share",
         "Illustrative policy test", "conc_change", "0.0%"),
    ]
    for lab, val, unit, note, nm, fmt in rows:
        put_input(ws, r, lab, val, unit, note, wb, nm, fmt)
        r += 1
    ws.cell(r, 1, "Bus share of bus + TEL fare revenue, FY2025/26")
    ws.cell(r, 2, "=Fund!B2/(Fund!B2+tel_rev/1000000)").number_format = "0.0%"
    ws.cell(r, 4, "Fund sheet (bus) and the rail contracting note")
    name(wb, "bus_share", f"Inputs!$B${r}")

    # ---------------------------------------------------------------- Fares
    ws = wb.create_sheet("Fares")
    cols = ["Band", "Fare km (ride-weighted mean of trunk rides)"]
    for tag in ("now", "Dec 2026"):
        for g in fm.GROUPS:
            cols += [f"{GROUP_LABEL[g]} trunk ({tag})", f"{GROUP_LABEL[g]} express ({tag})"]
    header(ws, 1, cols, [16, 14] + [11] * 16)
    for i in range(N):
        ws.cell(i + 2, 1, t.band[i])
        ws.cell(i + 2, 2, float(t.rep_km[i])).font = BLUE
        j = 3
        for tbl in (cur, new):
            for g in fm.GROUPS:
                ws.cell(i + 2, j, float(tbl[g]["trunk"][i])).font = BLUE
                ws.cell(i + 2, j + 1, float(tbl[g]["express"][i])).font = BLUE
                j += 2
    fr = N + 3
    ws.cell(fr, 1, "Feeder (flat)").font = BOLD
    j = 3
    for tbl in (cur, new):
        for g in fm.GROUPS:
            ws.cell(fr, j, float(tbl[g]["feeder"])).font = BLUE
            j += 2
    ws.cell(fr + 2, 1, "Cents per ride, card fares. Concession tables stop at 'Over 7.2 km'; longer bands repeat "
                       "that fare. Source: PTC FRE 2025 and FRE 2026 annexes (Sources sheet).")

    def fare_col(group, cls, tag):
        base = 3 + (0 if tag == "now" else 8) + 2 * fm.GROUPS.index(group)
        return L(base + (1 if cls == "express" else 0))

    def feeder_cell(group, tag):
        return f"Fares!${L(3 + (0 if tag == 'now' else 8) + 2 * fm.GROUPS.index(group))}${fr}"

    # ---------------------------------------------------------------- Rides
    ws = wb.create_sheet("Rides")
    header(ws, 1, ["Band", "Trunk rides / yr", "Feeder rides / yr", "Express rides / yr", "All"], [16, 16, 16, 16, 16])
    for i in range(N):
        ws.cell(i + 2, 1, t.band[i])
        for j, c in enumerate(["trunk", "feeder", "express"], 2):
            ws.cell(i + 2, j, float(t[c][i])).font = BLUE
            ws.cell(i + 2, j).number_format = "#,##0"
        ws.cell(i + 2, 5, f"=SUM(B{i + 2}:D{i + 2})").number_format = "#,##0"
    tr = N + 2
    ws.cell(tr, 1, "Total").font = BOLD
    for j in range(2, 6):
        ws.cell(tr, j, f"=SUM({L(j)}2:{L(j)}{N + 1})").number_format = "#,##0"
    name(wb, "rides_total", f"Rides!$E${tr}")
    ws.cell(tr + 2, 1, "Annual rides = mean of July and August 2026 rides per day by day type x 2026 day counts "
                       f"({stats['day_counts_2026']['WEEKDAY']} weekdays, "
                       f"{stats['day_counts_2026']['WEEKENDS/HOLIDAY']} weekend days and holidays). "
                       "Fare km = along-route distance on the cheapest service serving the stop pair.")

    # ---------------------------------------------------------------- Revenue
    ws = wb.create_sheet("Revenue")
    header(ws, 1, ["Band", "Blended trunk fare now (c)", "Blended express fare now (c)", "Revenue now (S$)",
                   "Adult-fare revenue now (S$)", "Blended trunk fare Dec 2026 (c)",
                   "Blended express fare Dec 2026 (c)", "Revenue Dec 2026 (S$)", "Adult trunk fare now (c)",
                   "Revenue with the concession test (S$)"],
           [16, 14, 14, 18, 18, 14, 14, 18, 12, 18])

    def blend(i, cls, tag, conc=False):
        parts = []
        for g in fm.GROUPS:
            ref = f"Fares!{fare_col(g, cls, tag)}{i + 2}"
            mult = "*(1+conc_change)" if (conc and g == "senior") else ""
            parts.append(f"share_{g}*{ref}{mult}")
        return "=" + "+".join(parts)

    def blend_feeder(tag, conc=False):
        return "+".join(f"share_{g}*{feeder_cell(g, tag)}" + ("*(1+conc_change)" if conc and g == "senior" else "")
                        for g in fm.GROUPS)

    for i in range(N):
        rr = i + 2
        ws.cell(rr, 1, t.band[i])
        ws.cell(rr, 2, blend(i, "trunk", "now"))
        ws.cell(rr, 3, blend(i, "express", "now"))
        feeder_now = f"({blend_feeder('now')})" if i == 0 else "0"
        ws.cell(rr, 4, f"=(Rides!B{rr}*B{rr}+Rides!C{rr}*{feeder_now}+Rides!D{rr}*C{rr})/100")
        a_feeder = f"{feeder_cell('adult', 'now')}" if i == 0 else "0"
        ws.cell(rr, 5, f"=(Rides!B{rr}*Fares!{fare_col('adult', 'trunk', 'now')}{rr}+Rides!C{rr}*{a_feeder}"
                       f"+Rides!D{rr}*Fares!{fare_col('adult', 'express', 'now')}{rr})/100")
        ws.cell(rr, 6, blend(i, "trunk", "dec"))
        ws.cell(rr, 7, blend(i, "express", "dec"))
        feeder_new = f"({blend_feeder('dec')})" if i == 0 else "0"
        ws.cell(rr, 8, f"=(Rides!B{rr}*F{rr}+Rides!C{rr}*{feeder_new}+Rides!D{rr}*G{rr})/100")
        ws.cell(rr, 9, f"=Fares!{fare_col('adult', 'trunk', 'now')}{rr}")
        fd = f"({blend_feeder('now', conc=True)})" if i == 0 else "0"
        ws.cell(rr, 10, f"=(Rides!B{rr}*({blend(i, 'trunk', 'now', conc=True)[1:]})+Rides!C{rr}*{fd}"
                        f"+Rides!D{rr}*({blend(i, 'express', 'now', conc=True)[1:]}))/100")
        ws.cell(rr, 10).number_format = "#,##0"
        for j in (4, 5, 8):
            ws.cell(rr, j).number_format = "#,##0"
    tr = N + 2
    ws.cell(tr, 1, "Total").font = BOLD
    for j in (4, 5, 8, 10):
        ws.cell(tr, j, f"=SUM({L(j)}2:{L(j)}{N + 1})").number_format = "#,##0"
    name(wb, "gross_now", f"Revenue!$D${tr}")
    name(wb, "adult_gross", f"Revenue!$E${tr}")
    name(wb, "gross_dec", f"Revenue!$H${tr}")
    k = tr + 2
    calc = [
        ("Mean adult fare per ride (cents)", "=adult_gross*100/rides_total", "mean_adult", "0.0"),
        ("Average transfer rebate F(a)+F(b)-F(a+b) (cents, Transfers sheet)", "=Transfers!B1", "rebate", "0.0"),
        ("Transfer factor (journey fares / ride fares)", "=1-transfer_share*rebate/mean_adult", "tfactor", "0.000"),
        ("Estimated fare revenue at current fares (S$/yr)", "=gross_now*tfactor", "rev_now", "#,##0"),
        ("Estimated fare revenue at Dec 2026 fares (S$/yr)", "=gross_dec*tfactor", "rev_dec", "#,##0"),
        ("Change from the Dec 2026 fare review", "=rev_dec/rev_now-1", "dec_change", "0.00%"),
        ("Extra revenue a year from the Dec 2026 fare review (S$)", "=rev_dec-rev_now", "dec_increase", "#,##0"),
        ("Revenue if every ride paid the adult fare, no transfers (S$/yr)", "=adult_gross", None, "#,##0"),
    ]
    for lab, f, nm, fmt in calc:
        ws.cell(k, 1, lab)
        ws.cell(k, 4, f).number_format = fmt
        if nm:
            name(wb, nm, f"Revenue!$D${k}")
        k += 1

    # ---------------------------------------------------------------- Transfers
    ws = wb.create_sheet("Transfers")
    ws.cell(1, 1, "Expected rebate per transfer (cents)").font = BOLD
    ws.cell(1, 2, f"=SUM(B4:{L(N + 1)}{N + 3})").number_format = "0.00"
    ws.cell(2, 1, "Cell (a, b) = P(a) P(b) [F(a) + F(b) - F(band of km_a + km_b)], adult trunk fares, P = share of "
                  "trunk rides in each band, km = Fares!B.")
    for i in range(N):
        ws.cell(3, i + 2, f"=Fares!$B${i + 2}")
        ws.cell(i + 4, 1, f"=Fares!$B${i + 2}")
    fa = fare_col("adult", "trunk", "now")
    for i in range(N):
        for j in range(N):
            pi, pj = f"Rides!$B${i + 2}/Rides!$B${N + 2}", f"Rides!$B${j + 2}/Rides!$B${N + 2}"
            comb = f"MIN({N - 1},MAX(0,ROUNDUP(ROUND(Fares!$B${i + 2}+Fares!$B${j + 2}-3.2,1)-0.000000001,0)))"
            ws.cell(i + 4, j + 2, f"={pi}*{pj}*(Fares!${fa}${i + 2}+Fares!${fa}${j + 2}"
                                  f"-INDEX(Fares!${fa}$2:${fa}${N + 1},{comb}+1))")

    # ---------------------------------------------------------------- Fund
    ws = wb.create_sheet("Fund")
    header(ws, 1, ["Financial year ended", "Fare revenue (S$m)", "Lease income (S$m)", "Other income (S$m)",
                   "Service fees and incentives (S$m)", "Total operating expenditure (S$m)",
                   "Deficit before grants (S$m)", "Operating grants (S$m)", "Fare revenue / operating cost",
                   "Deficit per ride (S$)"], [30, 14, 14, 14, 16, 16, 16, 14, 14, 14])
    for i, rec in enumerate(sorted(fund, key=lambda x: x["period"], reverse=True)):
        rr = i + 2
        ws.cell(rr, 1, rec["period"].replace("financial year ended ", ""))
        for j, key in enumerate(["fare_revenue", "bus_and_bus_related_lease_income", "other_operating_income",
                                 "service_fees_and_incentives", "total_operating_expenditure",
                                 "deficit_before_government_grants", "operating_grants"], 2):
            ws.cell(rr, j, rec.get(key) if rec.get(key) is not None else 0).font = BLUE
        ws.cell(rr, 9, f"=B{rr}/F{rr}").number_format = "0.0%"
    ws.cell(2, 10, "=-G2*1000000/(lta_daily*365)").number_format = "0.00"
    ws.cell(len(fund) + 3, 1, "Source: LTA audited financial statements, note on the Bus Contracting fund "
                              "(bus only; rail contracting is separate). Deficit per ride uses 2025 ridership.")
    name(wb, "fy25_fare", "Fund!$B$2")
    name(wb, "fy25_fees", "Fund!$E$2")
    name(wb, "fy25_opex", "Fund!$F$2")

    # ---------------------------------------------------------------- Reconciliation
    ws = wb.create_sheet("Reconciliation")
    header(ws, 1, ["Check", "Model", "Published", "Model / published - 1", "Note"], [52, 18, 18, 14, 90])
    rec_rows = [
        ("Annual bus rides", "=rides_total", "=lta_daily*365",
         "Model: July-August 2026 OD; published: LTA average daily bus ridership 2025 x 365"),
        ("Fare revenue, S$/yr (FY2025/26 restated to the Dec 2025 fares)", "=rev_now",
         "=fy25_fare*1000000+old_days/365*fre2025_inc*bus_share",
         "Published FY2025/26 bus fare revenue plus the part-year of the 2025 fare review it did not include"),
        ("Fare revenue, S$/yr (FY2025/26 as published)", "=rev_now", "=fy25_fare*1000000",
         "LTA audited Bus Contracting accounts, year ended 31 March 2026"),
        ("Extra revenue from the Dec 2026 fare review, S$/yr", "=dec_increase", "=fre2026_inc*bus_share",
         "PTC FRE 2026: bus + TEL increase, scaled to bus by the FY2025/26 fare revenue split"),
    ]
    for i, (lab, m, p, note) in enumerate(rec_rows, 2):
        ws.cell(i, 1, lab)
        ws.cell(i, 2, m).number_format = "#,##0"
        ws.cell(i, 3, p).number_format = "#,##0"
        ws.cell(i, 4, f"=B{i}/C{i}-1").number_format = "0.0%"
        ws.cell(i, 5, note)
    name(wb, "rec_rides_ratio", "Reconciliation!$D$2")
    name(wb, "pub_adj", "Reconciliation!$C$3")
    name(wb, "rec_rev_ratio", "Reconciliation!$D$3")
    name(wb, "rec_dec_ratio", "Reconciliation!$D$5")
    ws.cell(7, 1, "Implied yield: published revenue / revenue if every ride paid the adult fare")
    ws.cell(7, 2, "=pub_adj/adult_gross").number_format = "0.0%"
    name(wb, "implied_yield", "Reconciliation!$B$7")

    # ---------------------------------------------------------------- Contracts
    ws = wb.create_sheet("Contracts")
    header(ws, 1, ["Package", "Award date", "Operator", "Estimated total fee (S$)", "Years", "Services",
                   "Fee per year (S$)", "Fee per service-year (S$)", "Source"], [44, 12, 30, 18, 8, 10, 16, 16, 80])
    tendered = [c for c in sources["contracts"] if c.get("contract_type") == "tendered" and c.get("contract_value_sgd")]
    for i, c in enumerate(tendered, 2):
        ws.cell(i, 1, c["package"])
        ws.cell(i, 2, c["award_date"])
        ws.cell(i, 3, c["operator"])
        ws.cell(i, 4, c["contract_value_sgd"]).font = BLUE
        ws.cell(i, 5, c["contract_years"]).font = BLUE
        ws.cell(i, 6, c["n_services"]).font = BLUE
        ws.cell(i, 7, f"=D{i}/E{i}").number_format = "#,##0"
        ws.cell(i, 8, f"=G{i}/F{i}").number_format = "#,##0"
        s0 = c["sources"][0]
        ws.cell(i, 9, f"{s0['publisher']}: {s0['source_title']} {s0['url']}"
                      + (" (secondary source)" if any(s.get("secondary") for s in c["sources"]) else ""))
        ws.cell(i, 4).number_format = "#,##0"
    last = len(tendered) + 1
    k = last + 2
    ws.cell(k, 1, "All tendered awards: total fee per year / total services (S$ per service-year)")
    ws.cell(k, 7, f"=SUM(G2:G{last})/SUM(F2:F{last})").number_format = "#,##0"
    name(wb, "fee_service_year", f"Contracts!$G${k}")
    ws.cell(k + 1, 1, "System: FY2025/26 service fees / (scheduled weekday bus-km x 365), S$ per bus-km. Indicative: "
                      "the timetable estimate leaves out dead running and treats every day as a weekday")
    ws.cell(k + 1, 7, "=fy25_fees*1000000/(wd_bus_km*365)").number_format = "0.00"
    name(wb, "fee_per_bus_km", f"Contracts!$G${k + 1}")
    ws.cell(k + 2, 1, "Values are LTA's estimated total fee at award, excluding indexation, service variation and "
                      "incentives; terms exclude extension options.")

    # ---------------------------------------------------------------- Projection
    ws = wb.create_sheet("Projection")
    header(ws, 1, ["Financial year", "Fare level index", "Ridership index", "Fare revenue (S$m)",
                   "Other income (S$m)", "Operating cost (S$m)", "Deficit before grants (S$m)"],
           [16, 14, 14, 16, 14, 16, 18])
    ws.cell(2, 1, "FY2025/26 (actual)")
    ws.cell(2, 2, 1)
    ws.cell(2, 3, 1)
    ws.cell(2, 4, "=fy25_fare")
    ws.cell(2, 5, "=Fund!C2+Fund!D2")
    ws.cell(2, 6, "=fy25_opex")
    ws.cell(2, 7, "=D2+E2-F2")
    ws.cell(1, 8, "Full-year fare level").font = BOLD
    ws.cell(2, 8, "=1+fre2025_full")
    for y in range(1, 6):
        rr = 2 + y
        ws.cell(rr, 1, f"FY{2025 + y}/{(26 + y) % 100:02d}")
        # A late-December review covers 95 of the 365 days of the financial year in which it takes effect.
        step = "dec_change" if y == 1 else "later_fre"
        ws.cell(rr, 2, f"=H{rr - 1}*(1+95/365*{step})")
        ws.cell(rr, 8, f"=H{rr - 1}*(1+{step})")
        ws.cell(rr, 3, f"=C{rr - 1}*(1+g_rides)")
        ws.cell(rr, 4, f"=fy25_fare*B{rr}*C{rr}")
        ws.cell(rr, 5, "=E2")
        ws.cell(rr, 6, f"=F{rr - 1}*(1+g_cost)")
        ws.cell(rr, 7, f"=D{rr}+E{rr}-F{rr}")
        for jj in range(2, 9):
            ws.cell(rr, jj).number_format = "0.000" if jj in (2, 3, 8) else "#,##0"
    ws.cell(9, 1, "Fare level: FY2025/26 = 1. FY2026/27 adds the rest of the 2025 review (it applied to 95 of the 365 "
                  "days of FY2025/26) and 95 days of the Dec 2026 review; later years add the full-year effect of "
                  "the previous review and 95 days of the next.")
    ws.cell(10, 1, "Rest of the 2025 review (FY2025/26 restated / published - 1)")
    ws.cell(10, 2, "=pub_adj/(fy25_fare*1000000)-1").number_format = "0.00%"
    name(wb, "fre2025_full", "Projection!$B$10")
    name(wb, "deficit_fy2030", "Projection!$G$7")
    name(wb, "fare_fy2030", "Projection!$D$7")
    name(wb, "fare_fy2026", "Projection!$D$3")

    # ---------------------------------------------------------------- Scenarios
    ws = wb.create_sheet("Scenarios")
    ws.cell(1, 1, "Uniform fare change: estimated revenue (S$m/yr) by price elasticity of ridership").font = BOLD
    changes, elas = a["scenario_changes"], a["scenario_elasticities"]
    ws.cell(2, 1, "Fare change \\ elasticity")
    for j, e in enumerate(elas, 2):
        ws.cell(2, j, e).font = BLUE
    for i, ch in enumerate(changes, 3):
        ws.cell(i, 1, ch).font = BLUE
        ws.cell(i, 1).number_format = "0%"
        for j in range(2, len(elas) + 2):
            ws.cell(i, j, f"=rev_now*(1+$A{i})*(1+$A{i})^{L(j)}$2/1000000").number_format = "#,##0.0"
    r0 = len(changes) + 4
    ws.cell(r0, 1, f"Elasticities: {a['elasticity_note']}")
    r0 += 2
    ws.cell(r0, 1, "Concession test: change in senior / PwD fares (Inputs)").font = BOLD
    ws.cell(r0 + 1, 1, "Revenue with the concession change (S$/yr)")
    ws.cell(r0 + 1, 2, f"=Revenue!J{N + 2}*tfactor").number_format = "#,##0"
    ws.cell(r0 + 2, 1, "Change from current revenue (S$/yr)")
    ws.cell(r0 + 2, 2, f"=B{r0 + 1}-rev_now").number_format = "#,##0"
    name(wb, "conc_delta", f"Scenarios!$B${r0 + 2}")
    r0 += 4
    ws.cell(r0, 1, "Revenue-neutral flat adult fare (trunk rides)").font = BOLD
    ws.cell(r0 + 1, 1, "Flat fare (cents)")
    ws.cell(r0 + 1, 2, f"=SUMPRODUCT(Rides!B2:B{N + 1},Revenue!I2:I{N + 1})/Rides!B{N + 2}").number_format = "0.0"
    name(wb, "flat_fare", f"Scenarios!$B${r0 + 1}")
    ws.cell(r0 + 2, 1, "Share of trunk rides that would pay more than today")
    ws.cell(r0 + 2, 2, f"=SUMPRODUCT(Rides!B2:B{N + 1},--(Revenue!I2:I{N + 1}<flat_fare))/Rides!B{N + 2}")
    ws.cell(r0 + 2, 2).number_format = "0.0%"
    name(wb, "flat_more", f"Scenarios!$B${r0 + 2}")
    r0 += 4
    ws.cell(r0, 1, "Monthly fare spend of a regular adult commuter (S$)").font = BOLD
    header(ws, r0 + 1, ["Ride length", "Fare km", "Now", "Dec 2026", "Adult monthly pass",
                        "Now, % of 20th pct household income", "Dec 2026, % of that income"])
    for i, (lab, km) in enumerate([("Median ride", stats["fare_km_quantiles"]["0.5"]),
                                   ("90th percentile ride", stats["fare_km_quantiles"]["0.9"])], r0 + 2):
        ws.cell(i, 1, lab)
        ws.cell(i, 2, km).font = BLUE
        band = f"MIN({N - 1},MAX(0,ROUNDUP(ROUND(B{i}-3.2,1)-0.000000001,0)))+1"
        ws.cell(i, 3, f"=INDEX(Fares!${fare_col('adult', 'trunk', 'now')}$2:${fare_col('adult', 'trunk', 'now')}"
                      f"${N + 1},{band})*commuter_rides/100").number_format = "0.00"
        ws.cell(i, 4, f"=INDEX(Fares!${fare_col('adult', 'trunk', 'dec')}$2:${fare_col('adult', 'trunk', 'dec')}"
                      f"${N + 1},{band})*commuter_rides/100").number_format = "0.00"
        ws.cell(i, 5, "=adult_pass").number_format = "0.00"
        ws.cell(i, 6, f"=C{i}/income_p20").number_format = "0.00%"
        ws.cell(i, 7, f"=D{i}/income_p20").number_format = "0.00%"
    name(wb, "spend_median_now", f"Scenarios!$C${r0 + 2}")
    name(wb, "spend_p90_dec", f"Scenarios!$D${r0 + 3}")
    name(wb, "spend_p90_dec_income", f"Scenarios!$G${r0 + 3}")
    ws.column_dimensions["A"].width = 52

    # ---------------------------------------------------------------- Sources
    ws = wb.create_sheet("Sources")
    header(ws, 1, ["What", "Publisher", "Title", "URL", "Quoted text", "Accessed"], [26, 26, 50, 60, 70, 12])
    rr = 2
    for what, recs in _source_rows(sources):
        for s in recs:
            ws.cell(rr, 1, what)
            ws.cell(rr, 2, s.get("publisher"))
            ws.cell(rr, 3, s.get("source_title"))
            ws.cell(rr, 4, s.get("url"))
            ws.cell(rr, 5, s.get("quote"))
            ws.cell(rr, 6, s.get("accessed"))
            rr += 1
    for s in a.get("sources", []):
        for k, v in zip(range(1, 7), ["Assumption", s.get("publisher"), s.get("source_title"), s.get("url"),
                                      s.get("quote"), s.get("accessed")]):
            ws.cell(rr, k, v)
        rr += 1
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _source_rows(src):
    out = []
    for c in src["contracts"]:
        out.append((f"Contract: {c['package']}", c["sources"]))
    out.append(("Fare tables", src["fares"]["table_sources"]))
    out.append(("Monthly passes", src["fares"]["monthly_pass_sources"]))
    for f in src["fares"]["fare_review_exercises"]:
        out.append((f"Fare review {f['fre']}", f["sources"]))
    for r in src["revenue_cost"]:
        out.append((r.get("entity") or r.get("item"), [r]))
    for r in src["operations"]:
        out.append((r.get("item"), [r]))
    return out


def main():
    import finance
    table, od, km, stats = finance.build(fm.load_fares())
    (OUT / "finance_inputs.json").write_text(json.dumps(stats, indent=2), encoding="utf-8")
    path = build(stats=stats)
    print(f"wrote {path}")


if __name__ == "__main__":
    main()
